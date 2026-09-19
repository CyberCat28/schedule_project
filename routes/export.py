from flask import Blueprint, send_file
from firebase_admin import firestore
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

export_bp = Blueprint('export', __name__)

COLORS = {
    'header_bg': '132C54',   # тёмно-синяя шапка
    'group_bg': '536583',    # заголовки групп
    'day_bg': '132C54',      # строки дней
    'pair_bg': 'F4F6FA',     # столбец с парами
    'white': 'FFFFFF',
    'text_dark': '132C54',   # основной текст
    'text_light': '536583',  # приглушённый текст
    'border': 'D0D5DD'       # цвет границ
}

LESSON_TIMES = [
    {1: '9:00 - 10:35'},
    {2: '10:45 - 12:20'},
    {3: '13:05 - 14:40'},
    {4: '14:50 - 16:25'}
]

def side(style='thin', color='BDC3C7'):
    return Side(style=style, color=color)

def make_border(left='thin', right='thin', top='thin', bottom='thin'):
    return Border(left=side(left), right=side(right), top=side(top), bottom=side(bottom),)

thin_border = make_border()
medium_border = make_border('medium', 'medium', 'medium', 'medium')

def fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type='solid')

def is_light_color(hex_color):
    """ проверка, является ли цвет светлым """
    hex_color = str(hex_color).lstrip('#')
    if len(hex_color) != 6:
        return False

    try:
        r = int(hex_color[0:2], 16)
        g = int(hex_color[2:4], 16)
        b = int(hex_color[4:6], 16)
        brightness = (r * 299 + g * 587 + b * 114) / 1000
        return brightness > 155
    except ValueError:
        return False

def lesson_color(lesson, subject, teacher):
    return str( subject.get('color') or teacher.get('color') or '95A5A6').lstrip('#')

def lesson_text(lesson, teacher, subject):
    lines = [
        lesson.get('week_text', ''),
        lesson.get('type', ''),
        subject.get('short_name', ''),
        teacher.get('short_name', ''),
    ]

    if lesson.get('teacher_name'):
        lines.append(lesson['teacher_name'])

    if lesson.get('classroom'):
        lines.append(f"Каб. {lesson['classroom']}")

    # убирание пустых строк с сохранением порядка информации
    return '\n'.join(str(x) for x in lines if x)


def prepare_lesson(lesson, teacher_map, subject_map, pair):
    teacher = teacher_map.get(lesson.get('teacher_id'))
    subject = subject_map.get(lesson.get('subject_id'))

    if not teacher or not subject:
        return None

    weeks = []
    if lesson.get('week1_lesson') == pair:
        weeks.append('1н')
    if lesson.get('week2_lesson') == pair:
        weeks.append('2н')

    if not weeks:
        return None

    return {
        'id': lesson.get('id'),
        'teacher_short': teacher.get('short_name', ''),
        'subject_short': subject.get('short_name', ''),
        'type': lesson.get('lesson_type', ''),
        'classroom': lesson.get('classroom_id', ''),
        'weeks': ', '.join(weeks),
        'week_set': tuple(weeks),
        'color': lesson_color(lesson, subject, teacher),
    }


def unique_lessons_for_cell(lessons, group_id, day_idx, pair, teacher_map, subject_map):
    result = []

    for lesson in lessons:
        if lesson.get('group_id') != group_id:
            continue
        if lesson.get('day') != day_idx:
            continue

        prepared = prepare_lesson(lesson, teacher_map, subject_map, pair)
        if prepared is None:
            continue

        # одно и то же занятие, указанное сразу для 1-й и 2-й недели, считается одним блоком и занимает обе половины пары
        if not any(x['id'] == prepared['id'] for x in result):
            result.append(prepared)

    return result

@export_bp.route('/export/excel', methods=['GET'])
def export_to_excel():
    """ экспорт расписания в excel """
    try:
        db = firestore.client()

        groups = []
        for doc in db.collection('groups').stream():
            group = doc.to_dict()
            group['id'] = doc.id
            groups.append(group)

        teachers = []
        for doc in db.collection('teachers').stream():
            teacher = doc.to_dict()
            teacher['id'] = doc.id
            teachers.append(teacher)

        subjects = []
        for doc in db.collection('subjects').stream():
            subject = doc.to_dict()
            subject['id'] = doc.id
            subjects.append(subject)

        lessons = []
        for doc in db.collection('lessons').stream():
            lesson = doc.to_dict()
            lesson['id'] = doc.id
            lessons.append(lesson)

        teacher_map = {t['id']: t for t in teachers}
        subject_map = {s['id']: s for s in subjects}

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet(title='Расписание')

        days = ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница', 'Суббота',]

        # шапка
        a1 = ws.cell(row=1, column=1, value='Группа / Пара')
        a1.font = Font(bold=True, color=COLORS['white'], size=12)
        a1.fill = fill(COLORS['header_bg'])
        a1.alignment = Alignment(horizontal='center', vertical='center')
        a1.border = make_border('medium', 'medium', 'medium', 'thin')

        for col_idx, group in enumerate(groups, start=2):
            cell = ws.cell(row=1, column=col_idx, value=group.get('name', ''))
            cell.font = Font(bold=True, color=COLORS['white'], size=12)
            cell.fill = fill(COLORS['group_bg'])
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = make_border('thin', 'thin', 'thin', 'thin')

        current_row = 2

        # дни / пары
        for day_idx, day in enumerate(days):
            day_start = current_row

            # cтрока дня объединяется по всей ширине таблицы
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row,end_column=len(groups) + 1)
            day_cell = ws.cell(row=current_row, column=1, value=day)
            day_cell.font = Font(bold=True, color=COLORS['white'], size=12)
            day_cell.fill = fill(COLORS['day_bg'])
            day_cell.alignment = Alignment(horizontal='center', vertical='center',)
            day_cell.border = make_border('medium', 'thin', 'medium', 'thin')

            for col_idx in range(2, len(groups) + 2):
                c = ws.cell(row=current_row, column=col_idx)
                c.fill = fill(COLORS['day_bg'])
                c.border = make_border('thin', 'thin' if col_idx < len(groups) + 1 else 'medium', 'medium', 'thin')

            ws.row_dimensions[current_row].height = 24
            current_row += 1

            for pair in range(1, 5):
                # получение времени для этой пары
                lesson_time = ''
                for time_dict in LESSON_TIMES:
                    if pair in time_dict:
                        lesson_time = time_dict[pair]
                        break
                
                # собирание занятий по всем группам
                by_group = {}
                has_two_lessons = False

                for group in groups:
                    group_id = group['id']
                    cell_lessons = unique_lessons_for_cell(lessons, group_id, day_idx, pair, teacher_map, subject_map)
                    by_group[group_id] = cell_lessons
                    if len(cell_lessons) >= 2:
                        has_two_lessons = True

                # если есть две разные пары в одной ячейке, то создаются две строки и визуально разделяет их
                subrows = 2 if has_two_lessons else 1
                pair_start = current_row

                # номер пары и время объединяется по высоте блока.
                if subrows > 1:
                    ws.merge_cells(start_row=pair_start, start_column=1, end_row=pair_start + subrows - 1, end_column=1)

                pair_cell = ws.cell(row=pair_start, column=1, value=f'{pair} пара\n{lesson_time}')
                pair_cell.font = Font(bold=True, color=COLORS['text_light'], size=10)
                pair_cell.fill = fill(COLORS['pair_bg'])
                pair_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                pair_cell.border = make_border('medium', 'medium', 'thin', 'thin')

                for subrow in range(subrows):
                    ws.row_dimensions[pair_start + subrow].height = (75 if subrows == 2 else 90)

                # заполнение группы
                for col_idx, group in enumerate(groups, start=2):
                    group_id = group['id']
                    cell_lessons = by_group[group_id]

                    if subrows == 1:
                        target = ws.cell(row=pair_start, column=col_idx)

                        if cell_lessons:
                            lesson = cell_lessons[0]
                            target.value = (
                                f"{lesson['weeks']}\n"
                                f"\n{lesson['type']}\n"
                                f"{lesson['subject_short']}\n"
                                f"{lesson['teacher_short']}\n"
                            )
                            if lesson['classroom']:
                                target.value += f"\nКаб. {lesson['classroom']}"

                            target.fill = fill(lesson['color'])
                            target.font = Font(name='Calibri', size=11, bold=False, color=('555555' if is_light_color(lesson['color']) else '000000'))
                        else:
                            target.value = ''
                            target.fill = fill(COLORS['white'])
                            target.font = Font(size=11)

                        target.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                        target.border = make_border('thin', 'medium' if col_idx == len(groups) + 1 else 'thin', 'thin', 'thin')
                    else:
                        # две строки: если занятие одно, объединяется его на обе недели,а если два, то каждое получает свою строку
                        if len(cell_lessons) <= 1:
                            ws.merge_cells(start_row=pair_start, start_column=col_idx, end_row=pair_start + 1, end_column=col_idx)

                            target = ws.cell(row=pair_start, column=col_idx,)

                            if cell_lessons:
                                lesson = cell_lessons[0]
                                target.value = (
                                    f"{lesson['weeks']}\n"
                                    f"\n{lesson['type']}\n"
                                    f"{lesson['subject_short']}\n"
                                    f"{lesson['teacher_short']}\n"
                                )
                                if lesson['classroom']:
                                    target.value += f"\nКаб. {lesson['classroom']}"
                                target.fill = fill(lesson['color'])
                                target.font = Font(
                                    name='Calibri',
                                    size=11,
                                    color=('555555' if is_light_color(lesson['color']) else '000000'),
                                )
                            else:
                                target.value = ''
                                target.fill = fill(COLORS['white'])
                                target.font = Font(size=11)

                            target.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                            target.border = make_border('thin', 'medium' if col_idx == len(groups) + 1 else 'thin', 'thin', 'thin')

                            # нижняя граница объединённого блока
                            ws.cell(row=pair_start + 1, column=col_idx).border = make_border(
                                'thin', 'medium' if col_idx == len(groups) + 1 else 'thin', 'thin', 'thin',
                            )

                        else:
                            # сортировка (1-я неделя сверху, 2-я снизу)
                            def slot_key(x):
                                if '1н' in x['week_set']:
                                    return 0
                                return 1

                            cell_lessons = sorted(cell_lessons, key=slot_key)[:2]

                            for subrow, lesson in enumerate(cell_lessons):
                                target = ws.cell(row=pair_start + subrow, column=col_idx)
                                target.value = (
                                    f"{lesson['weeks']}\n"
                                    f"\n{lesson['type']}\n"
                                    f"{lesson['subject_short']}\n"
                                    f"{lesson['teacher_short']}\n"
                                )
                                if lesson['classroom']:
                                    target.value += f"\nКаб. {lesson['classroom']}"

                                target.fill = fill(lesson['color'])
                                target.font = Font(
                                    name='Calibri',
                                    size=11,
                                    color=(
                                        '555555' if is_light_color(lesson['color']) else '000000'
                                    ),
                                )
                                target.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                                target.border = make_border('thin', 'medium' if col_idx == len(groups) + 1 else 'thin', 'thin', 'thin')

                current_row += subrows

        # размеры
        ws.column_dimensions['A'].width = 18.0

        for col in range(2, len(groups) + 2):
            ws.column_dimensions[get_column_letter(col)].width = 25.7

        ws.row_dimensions[1].height = 30
        ws.freeze_panes = 'B2'

        ws.sheet_view.showGridLines = True

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype=('application/vnd.openxmlformats-officedocument.' 'spreadsheetml.sheet'),
            as_attachment=True,
            download_name='расписание.xlsx'
        )

    except Exception as e:
        return {'error': str(e)}, 500